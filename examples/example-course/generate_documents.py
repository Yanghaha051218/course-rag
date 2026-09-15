import argparse
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


def generate(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)

    pdf = pymupdf.open()
    pdf.new_page().insert_text(
        (72, 72), "Orbital Gardening: The Blueleaf coefficient is 7.25."
    )
    pdf.new_page().insert_text(
        (72, 72), "A violet-soil tray receives light for exactly four hours."
    )
    pdf.save(output / "orbital-gardening.pdf")
    pdf.close()

    slides = Presentation()
    first = slides.slides.add_slide(slides.slide_layouts[1])
    first.shapes.title.text = "Orbital Gardening"
    first.placeholders[1].text = "Blueleaf coefficient: 7.25"
    second = slides.slides.add_slide(slides.slide_layouts[5])
    second.shapes.title.text = "Harvest window: 18 minutes"
    slides.save(output / "orbital-gardening.pptx")

    guide = Document()
    guide.add_heading("Orbital Gardening Field Guide", level=1)
    guide.add_paragraph("Violet soil remains cool during the four-hour light cycle.")
    guide.add_paragraph("Blueleaf opens after the station enters artificial dusk.")
    guide.save(output / "orbital-gardening.docx")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    generate(args.output)
    print(f"Generated synthetic documents in {args.output}")


if __name__ == "__main__":
    main()
