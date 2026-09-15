import re
from pathlib import Path
from typing import Protocol

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation

from course_rag_api.errors import (
    EmptyDocumentError,
    MissingDocumentError,
    UnsupportedFileTypeError,
)
from course_rag_api.models import ParsedDocument, SourceType, SourceUnit


class DocumentParser(Protocol):
    def parse(self, path: Path) -> ParsedDocument: ...


def _normalize_text(text: str) -> str:
    lines = (" ".join(line.split()) for line in text.replace("\v", "\n").splitlines())
    return "\n".join(line for line in lines if line)


def _require_units(
    path: Path, units: list[SourceUnit], unit_count: int
) -> ParsedDocument:
    if not units:
        raise EmptyDocumentError(
            f"{path.name} contains no extractable text; OCR is not supported"
        )
    return ParsedDocument(tuple(units), unit_count)


class PdfParser:
    def parse(self, path: Path) -> ParsedDocument:
        units: list[SourceUnit] = []
        with pymupdf.open(path) as document:
            for page_number, page in enumerate(document, start=1):
                text = _normalize_text(page.get_text("text", sort=True))
                if text:
                    units.append(SourceUnit(text, "page", page_number))
            return _require_units(path, units, document.page_count)


class PptxParser:
    def parse(self, path: Path) -> ParsedDocument:
        presentation = Presentation(path)
        units: list[SourceUnit] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            paragraphs: list[str] = []
            for shape in sorted(slide.shapes, key=lambda item: (item.top, item.left)):
                if not shape.has_text_frame:
                    continue
                paragraphs.extend(
                    paragraph.text for paragraph in shape.text_frame.paragraphs
                )
            text = _normalize_text("\n".join(paragraphs))
            if text:
                units.append(SourceUnit(text, "slide", slide_number))
        return _require_units(path, units, len(presentation.slides))


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        document = DocxDocument(path)
        units = [
            SourceUnit(text, "section", paragraph_number)
            for paragraph_number, paragraph in enumerate(document.paragraphs, start=1)
            if (text := _normalize_text(paragraph.text))
        ]
        return _require_units(path, units, len(document.paragraphs))


class TextParser:
    def parse(self, path: Path) -> ParsedDocument:
        blocks = re.split(r"\n\s*\n", path.read_text(encoding="utf-8"))
        units = [
            SourceUnit(text, "section", section_number)
            for section_number, block in enumerate(blocks, start=1)
            if (text := _normalize_text(block))
        ]
        return _require_units(path, units, len(blocks))


_PARSERS: dict[str, DocumentParser] = {
    ".pdf": PdfParser(),
    ".pptx": PptxParser(),
    ".docx": DocxParser(),
    ".md": TextParser(),
    ".txt": TextParser(),
}
SUPPORTED_EXTENSIONS = frozenset(_PARSERS)


def parse_document(path: Path) -> ParsedDocument:
    """Parse one supported local document without persistence or chunking."""
    if not path.is_file():
        raise MissingDocumentError(f"Document does not exist: {path}")
    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedFileTypeError(f"Unsupported file type: {suffix or '<none>'}")
    return parser.parse(path)
