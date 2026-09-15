from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from course_rag_api.errors import (
    EmptyDocumentError,
    MissingDocumentError,
    UnsupportedFileTypeError,
)
from course_rag_api.parsers import parse_document


def test_pdf_parser_preserves_non_empty_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "orbital-gardening.pdf"
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "The Blueleaf coefficient is 7.25.")
    document.new_page()
    document.new_page().insert_text((72, 72), "Cafe orbit uses violet soil.")
    document.save(path)
    document.close()

    parsed = parse_document(path)

    assert parsed.unit_count == 3
    assert [(unit.source_type, unit.source_number) for unit in parsed.units] == [
        ("page", 1),
        ("page", 3),
    ]
    assert "Blueleaf coefficient is 7.25" in parsed.units[0].text


def test_pdf_without_extractable_text_fails(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        parse_document(path)


def test_pptx_parser_preserves_slide_numbers_and_order(tmp_path: Path) -> None:
    path = tmp_path / "greenhouse.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Orbital greenhouse"
    first.placeholders[1].text = "Blueleaf coefficient: 7.25"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    third = presentation.slides.add_slide(presentation.slide_layouts[5])
    third.shapes.title.text = "Harvest window: 18 minutes"
    presentation.save(path)

    parsed = parse_document(path)

    assert parsed.unit_count == 3
    assert [(unit.source_type, unit.source_number) for unit in parsed.units] == [
        ("slide", 1),
        ("slide", 3),
    ]
    assert parsed.units[0].text.index("Orbital greenhouse") < parsed.units[0].text.index(
        "Blueleaf coefficient"
    )


def test_docx_parser_uses_paragraph_units_not_pages(tmp_path: Path) -> None:
    path = tmp_path / "field-guide.docx"
    document = DocxDocument()
    document.add_paragraph("First observation: violet soil remains cool.")
    document.add_paragraph("")
    document.add_paragraph("Second observation: blueleaf opens after dusk.")
    document.save(path)

    parsed = parse_document(path)

    assert parsed.unit_count == 3
    assert [(unit.source_type, unit.source_number) for unit in parsed.units] == [
        ("section", 1),
        ("section", 3),
    ]


@pytest.mark.parametrize("suffix", [".md", ".txt"])
def test_text_parser_preserves_unicode_and_paragraphs(
    tmp_path: Path, suffix: str
) -> None:
    path = tmp_path / f"notes{suffix}"
    path.write_text("  Blueleaf   coefficient: 7.25.  \n\n 月光周期是四小时。\n")

    parsed = parse_document(path)

    assert parsed.unit_count == 2
    assert [unit.source_number for unit in parsed.units] == [1, 2]
    assert parsed.units[0].text == "Blueleaf coefficient: 7.25."
    assert parsed.units[1].text == "月光周期是四小时。"


def test_parser_rejects_missing_and_unsupported_files(tmp_path: Path) -> None:
    with pytest.raises(MissingDocumentError):
        parse_document(tmp_path / "missing.pdf")

    unsupported = tmp_path / "notes.csv"
    unsupported.write_text("not supported")
    with pytest.raises(UnsupportedFileTypeError, match=".csv"):
        parse_document(unsupported)
